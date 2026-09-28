"""FY working-paper summary, audit flags and XLSX export."""
import io, re
from collections import defaultdict
from datetime import date
from decimal import Decimal
from sqlalchemy.orm import Session
from ..models import Transaction, FinancialAccount, SourceDocument
from .rules import CATEGORIES, NEUTRAL, REVIEW, category_fits

LAKH = Decimal("100000")
MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]


def inr(v: Decimal) -> str:
    """Indian digit grouping: 12,34,567."""
    v = Decimal(v).quantize(Decimal("1"))
    s = str(abs(v))
    head, tail = s[:-3], s[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return ("-" if v < 0 else "") + "₹" + ",".join(groups + [tail]) if groups else ("-" if v < 0 else "") + "₹" + tail


def fy_transactions(db: Session, workspace_id: int, fy: str | None, account_id: int | None = None, purpose: str | None = None):
    q = db.query(Transaction, FinancialAccount).join(FinancialAccount, Transaction.account_id == FinancialAccount.id).filter(FinancialAccount.workspace_id == workspace_id)
    if fy:
        q = q.filter(Transaction.financial_year == fy)
    if account_id:
        q = q.filter(Transaction.account_id == account_id)
    if purpose:
        q = q.filter(Transaction.purpose == purpose)
    return q.order_by(Transaction.txn_date, Transaction.id).all()


def is_refund(t, a) -> bool:
    """A credit that gives back spending rather than bringing in money: any non-neutral credit on a card, a
    refund/reversal, or a credit in a spending category (e.g. a merchant credit filed under Dining)."""
    return t.category == "refund_reversal" or a.kind == "card" or not category_fits(t.category, False)


def totals(rows) -> dict:
    inflow = outflow = refunds = neutral = Decimal()
    for t, a in rows:
        if t.category in NEUTRAL:
            neutral += t.debit
            continue
        outflow += t.debit
        if is_refund(t, a):
            refunds += t.credit
        else:
            inflow += t.credit
    card = sum((t.debit for t, _ in rows if t.category == "card_settlement"), Decimal())
    return {"inflow": inflow, "outflow": outflow - refunds, "refunds": refunds, "neutral": neutral,
            "card_payments": card, "self_transfers": sum((t.debit for t, _ in rows if t.category == "own_transfer"), Decimal()),
            "card_emi": sum((t.debit for t, _ in rows if t.category == "card_emi"), Decimal())}


CASH = re.compile(r"\bcash\b", re.I)
# Card issuers as they appear in bank narrations for card bill payments.
ISSUERS = [(r"\bsbi ?card|sbicard", "SBI Card"), (r"amex|american express", "American Express"), (r"\bhdfc", "HDFC Bank"),
           (r"\bicici", "ICICI Bank"), (r"\baxis", "Axis Bank"), (r"kotak", "Kotak"), (r"idfc", "IDFC FIRST"), (r"\byes ?bank", "Yes Bank"),
           (r"indusind", "IndusInd"), (r"\brbl\b", "RBL"), (r"\bau\b.*(bank|card)", "AU Bank"), (r"standard chartered|\bscb\b", "Standard Chartered"),
           (r"\bhsbc\b", "HSBC"), (r"citi", "Citi"), (r"onecard|one card", "OneCard"), (r"federal", "Federal Bank"), (r"\bbob\b|baroda", "Bank of Baroda"),
           (r"\bsbi\b", "SBI Card")]
VIA = [(r"dreamplug|\bcred\b", "CRED"), (r"bbps|billdesk|bill ?pay", "BBPS"), (r"autopay|si-tad|si-mad|\bsi\b", "auto-debit")]
MONTH_NAMES = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def card_hint(narration: str) -> dict:
    """Which card a bank-side bill payment went to, as far as the narration says: issuer and last 4 digits."""
    low = narration.lower()
    via = next((name for pattern, name in VIA if re.search(pattern, low)), None)
    # Bank names inside UPI handles say whose rails carried the payment, not which card was paid:
    # CRED collects every card bill through cred.club@axisb, so "axis" there means nothing.
    low = re.sub(r"[\w.\-]+@[\w.\-]+", " ", low)
    # "POS 416021XXXXXX9685 ..." is the debit card that paid (e.g. through CRED), not the credit card being paid.
    paid_by_debit_card = re.match(r"\s*(pos\b|pos-|vps/|ips/)", low) is not None
    m = None if paid_by_debit_card else (re.search(r"(?:x{2,}|\*{2,}|\d{4}[x*]{2,})[x*\d]*?(\d{4})\b", low)
                                         or re.search(r"(?:credit ?card|\bcc\b|\bcard\b)\D{0,25}?(\d{4})\b", low))
    issuer = next((name for pattern, name in ISSUERS if re.search(pattern, low)), None)
    return {"issuer": issuer, "last4": m.group(1) if m else None, "via": via}


def _label(hint: dict) -> str:
    if hint["issuer"] or hint["last4"]:
        name = hint["issuer"] or "Credit"
        name = name if name.lower().endswith(("card", "express")) else f"{name} card"
        return name + (f" ••{hint['last4']}" if hint["last4"] else "")
    return f"Card not identified (paid via {hint['via']})" if hint["via"] else "Card not identified"


def _month(d) -> str:
    return f"{MONTH_NAMES[d.month - 1]} {d.year}"


def missing_card_statements(rows) -> list[dict]:
    """Bank-side card bill payments with no matching credit on an uploaded card statement, grouped by card, saying
    whether the card was never uploaded, a month is missing, or the statement is there but the payment isn't."""
    card_accounts = {a.id: a for _, a in rows if a.kind == "card"}
    coverage: dict[int, list[tuple]] = {}
    spans: dict[tuple, list] = {}
    for t, a in rows:
        if a.kind == "card":
            spans.setdefault((a.id, t.document_id), []).append(t.txn_date)
    for (acc_id, _), dates in spans.items():
        coverage.setdefault(acc_id, []).append((min(dates), max(dates)))
    for acc_id in coverage:
        coverage[acc_id].sort()
    from datetime import timedelta

    def find_account(h):
        cands = list(card_accounts.values())
        if h["last4"]:
            by_digits = [a for a in cands if h["last4"] in a.name]
            if by_digits:
                return by_digits[0]
            return None  # digits known but no uploaded card has them: a different card
        if h["issuer"]:
            by_issuer = [a for a in cands if h["issuer"].split()[0].lower() in a.name.lower()]
            return by_issuer[0] if len(by_issuer) == 1 else None
        return None

    groups: dict[tuple, dict] = {}
    for t, a in rows:
        if not (t.category == "card_settlement" and t.debit > 0 and a.kind == "bank" and not t.match_group):
            continue
        h = card_hint(t.narration)
        acc = find_account(h)
        if acc is None and not (h["issuer"] or h["last4"]):
            # CRED/BBPS payments don't say which card: group by month and name the cards with no statement for then.
            status, key = "not_identified", ("u", h["via"], _month(t.txn_date))
        elif acc is None:
            status, key = "not_uploaded", ("n", h["issuer"], h["last4"], None)
        else:
            covered = any(lo - timedelta(days=3) <= t.txn_date <= hi + timedelta(days=7) for lo, hi in coverage.get(acc.id, []))
            status, key = ("not_matched" if covered else "period_missing"), ("a", acc.id, "covered" if covered else "gap")
        covered = [f"{lo.strftime('%d %b %Y')} to {hi.strftime('%d %b %Y')}" for lo, hi in coverage.get(acc.id, [])] if acc else []
        g = groups.setdefault(key, {"card": acc.name if acc else _label(h), "account_id": acc.id if acc else None, "status": status,
                                    "count": 0, "amount": Decimal(), "months": [], "payments": [], "covered": covered, "via": h["via"]})
        g["count"] += 1
        g["amount"] += t.debit
        if _month(t.txn_date) not in g["months"]:
            g["months"].append(_month(t.txn_date))
        g["payments"].append({"id": t.id, "date": t.txn_date.isoformat(), "bank": a.name, "narration": t.narration, "amount": str(t.debit)})
        if status == "not_identified":
            # A payment's credit lands on the card statement whose period includes the payment date.
            missing = [c.name for c in card_accounts.values()
                       if not any(lo - timedelta(days=3) <= t.txn_date <= hi + timedelta(days=7) for lo, hi in coverage.get(c.id, []))]
            g.setdefault("uncovered", [])
            g["uncovered"] += [n for n in missing if n not in g["uncovered"]]
    order = {"not_identified": 0, "not_uploaded": 1, "period_missing": 2, "not_matched": 3}
    out = sorted(groups.values(), key=lambda g: (order[g["status"]], g["payments"][0]["date"] if g["status"] == "not_identified" else "", -g["amount"]))
    for g in out:
        g["amount"] = str(g["amount"])
        uncovered = g.get("uncovered", [])
        g["message"] = {
            "not_identified": (f"Paid via {g['via'] or 'a bill-pay app'}, which doesn't say which card, and "
                               "no uploaded card statement shows the payment or a total due it pays. "
                               + (f"Cards with no statement covering these dates: {', '.join(uncovered)}. Upload those statements, or those of any other card you paid."
                                  if uncovered else "Every uploaded card has a statement for these dates, so this probably paid a card you haven't uploaded at all.")),
            "not_uploaded": f"No statement uploaded for this card. Upload its statements for {', '.join(g['months'])} so its purchases are counted.",
            "period_missing": (f"This card is uploaded, but not the statement(s) covering {', '.join(g['months'])}. "
                               + (f"Statements uploaded for it cover: {'; '.join(g['covered'])}. " if g["covered"] else "None of its uploaded statements fall in this financial year. ")
                               + "Upload the missing ones so those purchases are counted."),
            "not_matched": "The statement for this period is uploaded, but no matching payment was found on it. Check the amount, or whether this payment went to a different card.",
        }[g["status"]]
    return out


def card_reconciliation(rows) -> dict:
    """Per credit card: purchases (already counted as expenses), refunds and bill payments, with how many
    payments were matched to a bank debit. Card bill payments are reported to the IT dept (AIS), so the
    total paid is shown for the CA to compare. Each payment is counted once even when both legs exist."""
    cards: dict[int, dict] = {}
    for t, a in rows:
        if a.kind != "card":
            continue
        c = cards.setdefault(a.id, {"account_id": a.id, "account": a.name, "purchases": Decimal(), "purchase_count": 0, "refunds": Decimal(),
                                    "payments": Decimal(), "payment_count": 0, "matched": Decimal(), "matched_count": 0,
                                    "unmatched": Decimal(), "unmatched_count": 0, "cash_payments": Decimal()})
        if t.category == "card_settlement" and t.credit > 0:
            c["payments"] += t.credit
            c["payment_count"] += 1
            if t.match_group:
                c["matched"] += t.credit
                c["matched_count"] += 1
            else:
                c["unmatched"] += t.credit
                c["unmatched_count"] += 1
            if CASH.search(t.narration):
                c["cash_payments"] += t.credit
        elif t.category == "refund_reversal" and t.credit > 0:
            c["refunds"] += t.credit
        elif t.debit > 0 and t.category not in NEUTRAL:
            c["purchases"] += t.debit
            c["purchase_count"] += 1
    # Bill payments made from an uploaded bank account to a card whose statement isn't uploaded.
    orphan = [(t, a) for t, a in rows if t.category == "card_settlement" and t.debit > 0 and a.kind == "bank" and not t.match_group]
    orphan_total = sum((t.debit for t, _ in orphan), Decimal())
    # Paid the total due of an uploaded statement; the credit is on a later statement that isn't uploaded.
    by_statement = [t for t, a in rows if t.status == "statement_paid" and a.kind == "bank"]
    statement_paid_total = sum((t.debit for t in by_statement), Decimal())
    card_list = sorted(cards.values(), key=lambda c: c["account"].lower())
    total_paid = sum((c["payments"] for c in card_list), Decimal()) + orphan_total + statement_paid_total
    cash_total = sum((c["cash_payments"] for c in card_list), Decimal())
    as_str = lambda d: {k: (str(v) if isinstance(v, Decimal) else v) for k, v in d.items()}
    return {
        "cards": [as_str(c) for c in card_list],
        "missing_statements": missing_card_statements(rows),
        "unmatched_bank_payments": [{"id": t.id, "date": t.txn_date.isoformat(), "account": a.name, "narration": t.narration, "amount": str(t.debit)} for t, a in orphan],
        "unmatched_bank_total": str(orphan_total),
        "statement_paid_count": len(by_statement), "statement_paid_total": str(statement_paid_total),
        "total_paid": str(total_paid),
        "total_purchases": str(sum((c["purchases"] for c in card_list), Decimal())),
        "cash_paid": str(cash_total),
        "sft_reportable": total_paid >= 10 * LAKH or cash_total >= LAKH,
    }


BAD_READ = ("Doesn't add up", "Only one transaction was read", "Running balance breaks")
NO_ROWS = "No rows were read for this month, though the statement runs through it"


def statement_coverage(db: Session, workspace_id: int, fy: str | None) -> dict:
    """Which statements are uploaded, per account and month of the financial year.

    A card statement belongs to the month its cycle ends (its last transaction, about the statement date); a bank
    statement covers every month from its first to its last transaction. A month is "warn" when its statement
    didn't read cleanly. `unexplained` counts, per month, bank card-bill payments that no uploaded card explains."""
    from sqlalchemy import func
    start = int(fy[:4]) if fy else date.today().year - (1 if date.today().month < 4 else 0)
    months = [f"{start + (m < 4)}-{m:02d}" for m in [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]]
    accounts = db.query(FinancialAccount).filter_by(workspace_id=workspace_id).all()
    spans = {doc_id: (lo, hi, n) for doc_id, lo, hi, n in
             db.query(Transaction.document_id, func.min(Transaction.txn_date), func.max(Transaction.txn_date), func.count(Transaction.id))
             .group_by(Transaction.document_id)}
    doc_months = defaultdict(set)
    for doc_id, d in db.query(Transaction.document_id, Transaction.txn_date).join(FinancialAccount).filter(FinancialAccount.workspace_id == workspace_id):
        doc_months[doc_id].add(f"{d.year}-{d.month:02d}")
    out = []
    for a in sorted(accounts, key=lambda a: (a.kind != "card", a.name.lower())):
        cells: dict[str, dict] = {}
        for d in db.query(SourceDocument).filter_by(workspace_id=workspace_id, account_id=a.id):
            if d.id not in spans:
                continue
            lo, hi, n = spans[d.id]
            bad = any(w in (d.warnings or "") for w in BAD_READ)
            if a.kind == "card":
                covered = [f"{hi.year}-{hi.month:02d}"]
            else:
                covered, y, m = [], lo.year, lo.month
                while (y, m) <= (hi.year, hi.month):
                    covered.append(f"{y}-{m:02d}")
                    y, m = (y + 1, 1) if m == 12 else (y, m + 1)
            for key in covered:
                c = cells.setdefault(key, {"status": "ok", "statements": []})
                # A bank statement spanning a month with no rows in it was most likely not read in full.
                empty = a.kind == "bank" and key not in doc_months[d.id]
                c["statements"].append({"id": d.id, "from": lo.isoformat(), "to": hi.isoformat(), "rows": n, "total_due": str(d.total_due) if d.total_due is not None else None,
                                        "problem": NO_ROWS if empty else next((w for w in BAD_READ if w in (d.warnings or "")), None)})
                if bad or empty:
                    c["status"] = "warn"
        out.append({"account_id": a.id, "account": a.name, "kind": a.kind,
                    "months": {m: cells.get(m, {"status": "missing", "statements": []}) for m in months},
                    "uploaded": sum(1 for m in months if m in cells)})
    unexplained = defaultdict(int)
    for t, a in fy_transactions(db, workspace_id, fy):
        if a.kind == "bank" and t.category == "card_settlement" and t.debit > 0 and not t.match_group:
            unexplained[f"{t.txn_date.year}-{t.txn_date.month:02d}"] += 1
    return {"months": months, "accounts": out, "unexplained": {m: unexplained.get(m, 0) for m in months}}


def card_breakdown(rows) -> dict:
    """Per credit card: what was paid to it, then what its statements say that money went on.

    paid: bill payments on the card's statements, plus bank payments that paid one of its statements' total due
    (their credit is on a later, missing statement). From the statements: purchases by category, charges/fees,
    refunds, EMI instalments (a purchase converted to EMI is counted once, as the purchase), and net spend.
    Bank card-bill payments tied to no card are listed on their own so paid totals still add up."""
    docs_account = {t.document_id: a.id for t, a in rows if a.kind == "card"}
    cards: dict[int, dict] = {}

    def card(a):
        return cards.setdefault(a.id, {"account_id": a.id, "account": a.name, "paid": Decimal(), "payments": 0, "purchases": Decimal(),
                                       "charges": Decimal(), "refunds": Decimal(), "emi": Decimal(), "categories": defaultdict(Decimal),
                                       "from": None, "to": None})
    unassigned = {"count": 0, "amount": Decimal()}
    for t, a in rows:
        if a.kind == "card":
            c = card(a)
            c["from"] = min(c["from"] or t.txn_date, t.txn_date)
            c["to"] = max(c["to"] or t.txn_date, t.txn_date)
            if t.category == "card_settlement" and t.credit > 0:
                c["paid"] += t.credit
                c["payments"] += 1
            elif t.category == "card_emi":
                c["emi"] += t.debit
            elif t.credit > 0:
                c["refunds"] += t.credit
            elif t.category == "bank_charges":
                c["charges"] += t.debit
            elif t.debit > 0:
                c["purchases"] += t.debit
                c["categories"][t.category] += t.debit
        elif t.category == "card_settlement" and t.debit > 0 and not (t.match_group and not t.match_group.startswith("doc:")):
            # Not linked to a card-side credit (that one is counted on the card): either it paid an uploaded
            # statement ("doc:<id>") or no card is known.
            doc_id = int(t.match_group[4:]) if t.match_group else None
            if doc_id in docs_account:
                c = cards.get(docs_account[doc_id])
                if c is not None:
                    c["paid"] += t.debit
                    c["payments"] += 1
                    continue
            unassigned["count"] += 1
            unassigned["amount"] += t.debit
    out = []
    for c in sorted(cards.values(), key=lambda c: c["account"].lower()):
        cats = sorted(c.pop("categories").items(), key=lambda kv: -kv[1])
        spend = c["purchases"] + c["charges"] - c["refunds"]
        out.append({**{k: (str(v) if isinstance(v, Decimal) else v) for k, v in c.items()},
                    "from": c["from"].isoformat() if c["from"] else None, "to": c["to"].isoformat() if c["to"] else None,
                    "net_spend": str(spend),
                    "categories": [{"category": k, "label": CATEGORIES.get(k, (k,))[0], "amount": str(v)} for k, v in cats]})
    return {"cards": out, "paid_total": str(sum((Decimal(c["paid"]) for c in out), Decimal()) + unassigned["amount"]),
            "spend_total": str(sum((Decimal(c["net_spend"]) for c in out), Decimal())),
            "unassigned": {"count": unassigned["count"], "amount": str(unassigned["amount"])}}


EXCLUDED_LABELS = {
    ("card_settlement", "bank"): "Credit card bills paid from the bank",
    ("card_settlement", "card"): "Bill payments received on the cards",
    ("own_transfer", "bank"): "Transfers between your own accounts",
    ("own_transfer", "card"): "Transfers between your own accounts",
    ("card_emi", "card"): "Card EMI conversions and instalments",
    ("card_emi", "bank"): "Card EMI conversions and instalments",
}


def books(rows) -> dict:
    """One date-ordered ledger across every bank account and credit card: each purchase, receipt and refund once.

    Both legs of a card bill payment (the bank debit and the payment credit on the card) are left out, since the
    card purchases themselves are in the ledger; so are self-transfers and card EMI conversions. What was left out
    is totalled in `excluded`, so the ledger can be checked against the statements. Bill payments to cards whose
    statement isn't uploaded are listed under `gaps`: those purchases are missing from the ledger until it is.
    """
    entries, excluded = [], {}
    money_in = money_out = refunds = running = Decimal()
    for t, a in rows:
        if t.category in NEUTRAL:
            label = EXCLUDED_LABELS.get((t.category, a.kind), CATEGORIES[t.category][0])
            e = excluded.setdefault(label, {"label": label, "count": 0, "debit": Decimal(), "credit": Decimal(), "matched": 0})
            e["count"] += 1
            e["debit"] += t.debit
            e["credit"] += t.credit
            e["matched"] += bool(t.match_group)
            continue
        refund = t.credit > 0 and is_refund(t, a)
        money_out += t.debit
        if refund:
            refunds += t.credit
        else:
            money_in += t.credit
        running += t.credit - t.debit
        label, group, _ = CATEGORIES.get(t.category, (t.category, "review", ""))
        entries.append({"id": t.id, "date": t.txn_date.isoformat(), "account": a.name, "account_kind": a.kind, "narration": t.narration,
                        "category": t.category, "category_label": label, "group": group, "out": str(t.debit), "in": str(t.credit),
                        "refund": refund, "running": str(running), "status": t.status, "note": t.note, "purpose": t.purpose})
    gaps = missing_card_statements(rows)
    return {"entries": entries,
            "totals": {"money_in": str(money_in), "money_out": str(money_out), "refunds": str(refunds),
                       "net_spend": str(money_out - refunds), "net": str(running), "count": len(entries)},
            "excluded": [{**e, "debit": str(e["debit"]), "credit": str(e["credit"])} for e in excluded.values()],
            "gaps": gaps, "gap_total": str(sum((Decimal(g["amount"]) for g in gaps), Decimal())),
            "cards": card_breakdown(rows)}


def purpose_split(rows) -> dict:
    """Money in/out and counts by business / personal / unknown (neutral flows excluded)."""
    out = {}
    for p in ("business", "personal", "unknown"):
        sel = [(t, a) for t, a in rows if t.purpose == p]
        tot = totals(sel)
        out[p] = {"count": len(sel), "inflow": str(tot["inflow"]), "outflow": str(tot["outflow"])}
    return out


def purpose_flags(rows, purpose: str | None) -> list[dict]:
    """In a business/personal report, say what was left out so nothing silently disappears."""
    if not purpose:
        return []
    unknown = [t for t, _ in rows if t.purpose == "unknown"]
    other = "personal" if purpose == "business" else "business"
    excluded = [t for t, _ in rows if t.purpose == other]
    out = []
    if unknown:
        out.append({"level": "warning", "title": f"{len(unknown)} transaction(s) have no business/personal purpose yet",
                    "detail": f"{inr(sum(t.debit + t.credit for t in unknown))} is left out of this {purpose} report until you choose. "
                              "Open Transactions → Purpose: Unknown, or set what each account is used for in Settings."})
    if excluded:
        out.append({"level": "info", "title": f"{len(excluded)} {other} transaction(s) excluded",
                    "detail": f"{inr(sum(t.debit for t in excluded))} out and {inr(sum(t.credit for t in excluded))} in are marked {other} and not part of this report."})
    return out


def account_summary(rows) -> list[dict]:
    """One line per account: the same totals as the overall report, split by account."""
    by_account: dict[int, list] = defaultdict(list)
    for t, a in rows:
        by_account[a.id].append((t, a))
    out = []
    for acc_rows in by_account.values():
        a = acc_rows[0][1]
        tot = totals(acc_rows)
        dates = [t.txn_date for t, _ in acc_rows]
        out.append({"account_id": a.id, "account": a.name, "kind": a.kind, "count": len(acc_rows),
                    "inflow": str(tot["inflow"]), "outflow": str(tot["outflow"]), "refunds": str(tot["refunds"]), "neutral": str(tot["neutral"]),
                    "exceptions": sum(t.status in {"needs_review", "unmatched", "ambiguous"} for t, _ in acc_rows),
                    "from": min(dates).isoformat(), "to": max(dates).isoformat()})
    return sorted(out, key=lambda r: (r["kind"] != "bank", r["account"].lower()))


def category_summary(rows) -> list[dict]:
    agg = defaultdict(lambda: {"count": 0, "debit": Decimal(), "credit": Decimal()})
    for t, _ in rows:
        a = agg[t.category]
        a["count"] += 1
        a["debit"] += t.debit
        a["credit"] += t.credit
    out = []
    for key, a in agg.items():
        label, group, hint = CATEGORIES.get(key, (key, "review", ""))
        out.append({"category": key, "label": label, "group": group, "itr_hint": hint, "count": a["count"], "debit": str(a["debit"]), "credit": str(a["credit"])})
    order = ["income", "tax", "deduction_hint", "investment", "review", "expense", "adjustment", "neutral"]
    return sorted(out, key=lambda r: (order.index(r["group"]) if r["group"] in order else 99, -Decimal(r["debit"]) - Decimal(r["credit"])))


def flags(rows, fy: str | None) -> list[dict]:
    out = []

    def add(level, title, detail):
        out.append({"level": level, "title": title, "detail": detail})

    by_status = defaultdict(list)
    for t, a in rows:
        by_status[(t.category, t.status, a.kind, t.debit > 0)].append(t)
    bank_card_payments = [t for (c, s, k, d), ts in by_status.items() if c == "card_settlement" and s == "unmatched" and d for t in ts]
    card_payments_in = [t for (c, s, k, d), ts in by_status.items() if c == "card_settlement" and s == "unmatched" and not d for t in ts]
    transfers = [t for (c, s, k, d), ts in by_status.items() if c == "own_transfer" and s == "unmatched" for t in ts]
    ambiguous = [t for t, _ in rows if t.status == "ambiguous"]
    review = [t for t, _ in rows if t.status == "needs_review"]

    for g in missing_card_statements(rows):
        title = {"not_identified": f"{g['card']}, {', '.join(g['months'])}: no card statement shows it",
                 "not_uploaded": f"{g['card']}: statement not uploaded",
                 "period_missing": f"{g['card']}: statement missing for {', '.join(g['months'])}",
                 "not_matched": f"{g['card']}: payment not found on the uploaded statement"}[g["status"]]
        add("warning", title, f"{g['count']} bill payment(s) of {inr(Decimal(g['amount']))} from your bank ({', '.join(g['months'])}). {g['message']}")
    if card_payments_in:
        add("info", f"{len(card_payments_in)} card payment(s) received with no paying bank statement",
            f"{inr(sum(t.credit for t in card_payments_in))} of card bills were paid from an account that isn't uploaded. Upload that bank statement if it is yours.")
    if transfers:
        add("info", f"{len(transfers)} self-transfer(s) without the other side",
            "These look like transfers between your own accounts but the other account's statement isn't uploaded. Upload it, or re-categorise if the money went to someone else.")
    if ambiguous:
        add("warning", f"{len(ambiguous)} payment(s) could match more than one counterpart",
            "Several same-amount payments fall on nearby dates. Check these in Reconciliation; nothing was auto-linked.")
    if review:
        add("warning", f"{len(review)} transaction(s) need a category",
            f"{inr(sum(t.debit + t.credit for t in review))} of UPI/NEFT/unrecognised entries. Open Transactions → 'Needs review' and assign categories.")

    big = [t for t, _ in rows if t.credit >= 2 * LAKH and (t.category in REVIEW or t.category in {"business_receipt"})]
    if big:
        add("warning", f"{len(big)} large credit(s) of ₹2 lakh or more with no clear source",
            "The tax department can ask you to explain large credits (sec. 68/69). Note the source of each — e.g. loan, gift from relative, sale of asset. "
            + "; ".join(f"{t.txn_date.isoformat()} {inr(t.credit)}" for t in big[:5]) + ("…" if len(big) > 5 else ""))

    cash_dep = sum((t.credit for t, a in rows if t.category == "cash_deposit"), Decimal())
    if cash_dep >= 10 * LAKH:
        add("warning", f"Cash deposits of {inr(cash_dep)} this year",
            "Cash deposits of ₹10 lakh or more in a year are reported by banks to the IT dept (SFT) and appear in your AIS. Keep evidence of the source.")
    elif cash_dep > 0:
        add("info", f"Cash deposits of {inr(cash_dep)}", "Below the ₹10 lakh SFT reporting threshold, but keep a note of the source.")

    card_paid = Decimal(card_reconciliation(rows)["total_paid"])
    if card_paid >= 10 * LAKH:
        add("info", f"Credit card bills paid: {inr(card_paid)}", "Card bill payments of ₹10 lakh or more in a year are reported in AIS (SFT). Make sure your declared income supports this spend.")

    interest = sum((t.credit for t, _ in rows if t.category == "interest_income"), Decimal())
    if interest > 0:
        add("info", f"Interest received: {inr(interest)}", "Declare under 'Income from other sources' and compare with AIS/26AS — FD interest is often missed.")
    dividend = sum((t.credit for t, _ in rows if t.category == "dividend"), Decimal())
    if dividend > 0:
        add("info", f"Dividends received: {inr(dividend)}", "Dividends are taxable at slab rates. Match with AIS and claim any TDS deducted.")
    salary = sum((t.credit for t, _ in rows if t.category == "salary"), Decimal())
    if salary > 0:
        add("info", f"Salary credits: {inr(salary)} (net of deductions)", "Bank credits are take-home pay. Use gross salary and TDS from Form 16 for the return.")

    if fy:
        start = int(fy[:4])
        months_in_fy = [(start if m >= 4 else start + 1, m) for m in [4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3]]
        today = date.today()
        by_account = defaultdict(set)
        names, spans = {}, {}
        for t, a in rows:
            by_account[a.id].add((t.txn_date.year, t.txn_date.month))
            names[a.id] = a.name
            if a.kind == "bank":
                lo, hi, _ = spans.get(t.document_id, (t.txn_date, t.txn_date, a.id))
                spans[t.document_id] = (min(lo, t.txn_date), max(hi, t.txn_date), a.id)
        for acc_id, seen in by_account.items():
            gaps = [ym for ym in months_in_fy if ym not in seen and date(ym[0], ym[1], 1) <= today]
            missing = [MONTHS[months_in_fy.index(ym)] for ym in gaps]
            if missing and len(missing) < 12:
                # Inside the dates of an uploaded statement: that file was read only in part, it isn't missing.
                inside = [MONTHS[months_in_fy.index(ym)] for ym in gaps
                          if any(acc == acc_id and (lo.year, lo.month) < ym < (hi.year, hi.month) for lo, hi, acc in spans.values())]
                if inside:
                    add("warning", f"{names[acc_id]}: no transactions read for {', '.join(inside)}",
                        "An uploaded statement runs through these months, but no rows were read for them, so part of the file was "
                        "probably not read. Check that statement's import warnings under Statement coverage, then delete and re-upload it.")
                outside = [m for m in missing if m not in inside]
                if outside:
                    add("info", f"{names[acc_id]}: no transactions in {', '.join(outside)}",
                        "This may be a missing statement period. Upload it if the account was active.")
    if not rows:
        add("info", "No statements uploaded for this year", "Upload bank and credit card statements to begin.")
    return out


def export_xlsx(db: Session, workspace_id: int, fy: str | None, account_id: int | None = None, purpose: str | None = None,
                business_mode: bool = False) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    all_rows = fy_transactions(db, workspace_id, fy, account_id)
    rows = [(t, a) for t, a in all_rows if t.purpose == purpose] if purpose else all_rows
    docs = {d.id: d.filename for d in db.query(SourceDocument).filter_by(workspace_id=workspace_id)}
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="D9EEE4")

    ws = wb.active
    ws.title = "Summary"
    tot = totals(rows)
    scope = f" — {all_rows[0][1].name}" if account_id and all_rows else ""
    kind = f"{purpose.capitalize()} working paper" if purpose else "working paper"
    ws.append([f"LedgerVault {kind} — FY {fy or 'all years'}{scope}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Provisional. Prepared from bank/card statements for review by a Chartered Accountant; not a tax computation."])
    if purpose:
        split = purpose_split(all_rows)
        ws.append([f"Only transactions marked {purpose}. Excluded: {split['personal' if purpose == 'business' else 'business']['count']} "
                   f"{'personal' if purpose == 'business' else 'business'} and {split['unknown']['count']} with no purpose yet."])
    ws.append([])
    for label, key in [("Money in (excl. transfers & refunds)", "inflow"), ("Money out (net of refunds, excl. card bill payments & transfers)", "outflow"),
                       ("Refunds / reversals", "refunds"), ("Credit card bill payments (not an expense — see 'Credit cards' sheet)", "card_payments"),
                       ("Transfers between own accounts (not income or expense)", "self_transfers")]:
        ws.append([label, float(tot[key])])
    ws.append([])
    ws.append(["Category", "Group", "Count", "Debit (₹)", "Credit (₹)", "ITR note"])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, head_fill
    for r in category_summary(rows):
        if r["group"] == "neutral":
            continue  # card bill payments and self-transfers are listed above and reconciled on their own sheet
        ws.append([r["label"], r["group"].replace("_", " "), r["count"], float(r["debit"]), float(r["credit"]), r["itr_hint"]])
    ws.append(["Credit card bill payments and transfers between own accounts are not income or expenses; see the lines above and the 'Credit cards' sheet."])
    for col, width in zip("ABCDEF", [48, 16, 8, 16, 16, 80]):
        ws.column_dimensions[col].width = width

    bk = books(rows)
    bs = wb.create_sheet("Books")
    bs.append(["Books — every bank and credit card transaction once, in date order"])
    bs["A1"].font = Font(bold=True, size=13)
    bs.append(["Credit card bill payments (the bank debit and the card's payment credit), transfers between own accounts and card EMI "
               "conversions are left out: the card purchases themselves are listed. What was left out is totalled at the bottom."])
    bs.append([])
    bs.append(["Date", "Account", "Type", "Narration", "Category", "Money out (₹)", "Money in (₹)", "Refund?", "Running net (₹)", "Status", "Note"])
    for c in bs[bs.max_row]:
        c.font, c.fill = bold, head_fill
    for e in bk["entries"]:
        bs.append([date.fromisoformat(e["date"]), e["account"], "Card" if e["account_kind"] == "card" else "Bank", e["narration"], e["category_label"],
                   float(e["out"]) or None, float(e["in"]) or None, "refund" if e["refund"] else "", float(e["running"]),
                   e["status"].replace("_", " "), e["note"] or ""])
    bs.append([])
    for label, key in [("Money out", "money_out"), ("Money in (excl. refunds)", "money_in"), ("Refunds / reversals", "refunds"),
                       ("Net spend (money out − refunds)", "net_spend"), ("Net (in − out)", "net")]:
        bs.append([label, None, None, None, None, float(bk["totals"][key])])
        bs[bs.max_row][0].font = bold
    bs.append([])
    bs.append(["Left out of the books", None, None, "Count", "Matched to the other leg", "Debits (₹)", "Credits (₹)"])
    for c in bs[bs.max_row]:
        c.font, c.fill = bold, head_fill
    for e in bk["excluded"]:
        bs.append([e["label"], None, None, e["count"], e["matched"], float(e["debit"]), float(e["credit"])])
    cb = bk["cards"]
    if cb["cards"]:
        bs.append([])
        bs.append(["Credit cards: what was paid to each card, and what its statements say it went on"])
        bs[bs.max_row][0].font = bold
        bs.append(["Card", "Statements cover", None, "Paid to card (₹)", "Purchases (₹)", "Fees & charges (₹)", "Refunds (₹)", "Net spend (₹)", "EMI instalments (₹)"])
        for c in bs[bs.max_row]:
            c.font, c.fill = bold, head_fill
        for c in cb["cards"]:
            bs.append([c["account"], f"{c['from']} to {c['to']}", None, float(c["paid"]), float(c["purchases"]), float(c["charges"]),
                       float(c["refunds"]), float(c["net_spend"]), float(c["emi"])])
            for cat in c["categories"]:
                bs.append([f"    {cat['label']}", None, None, None, float(cat["amount"])])
        if cb["unassigned"]["count"]:
            bs.append([f"Card bill payments not tied to any uploaded card ({cb['unassigned']['count']})", None, None, float(cb["unassigned"]["amount"])])
        bs.append(["Total", None, None, float(cb["paid_total"]), None, None, None, float(cb["spend_total"])])
        bs[bs.max_row][0].font = bold
    if bk["gaps"]:
        bs.append([])
        bs.append([f"Not in the books yet: {inr(Decimal(bk['gap_total']))} of card bills paid to cards whose statements aren't uploaded "
                   "(their purchases are missing). See the 'Credit cards' sheet."])
        bs[bs.max_row][0].font = Font(bold=True, color="B00020")
    bs.freeze_panes = "A5"
    for col, width in zip("ABCDEFGHIJK", [12, 26, 7, 60, 28, 15, 15, 9, 16, 16, 30]):
        bs.column_dimensions[col].width = width

    if not account_id:
        ba = wb.create_sheet("By account")
        ba.append(["Account", "Type", "Period", "Transactions", "Money in (₹)", "Money out, net (₹)", "Refunds (₹)", "Neutral (₹)", "Needs attention"])
        for c in ba[1]:
            c.font, c.fill = bold, head_fill
        for r in account_summary(rows):
            ba.append([r["account"], "Credit card" if r["kind"] == "card" else "Bank", f"{r['from']} to {r['to']}", r["count"],
                       float(r["inflow"]), float(r["outflow"]), float(r["refunds"]), float(r["neutral"]), r["exceptions"]])
        # Category × account matrix, so the CA can see where each category came from.
        accounts = [(r["account_id"], r["account"]) for r in account_summary(rows)]
        ba.append([])
        ba.append(["Category (net: in − out)"] + [name for _, name in accounts] + ["Total"])
        for c in ba[ba.max_row]:
            c.font, c.fill = bold, head_fill
        matrix: dict[str, dict[int, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
        for t, a in rows:
            matrix[t.category][a.id] += t.credit - t.debit
        for cat in sorted(matrix, key=lambda k: CATEGORIES.get(k, (k,))[0]):
            vals = [matrix[cat].get(acc_id, Decimal()) for acc_id, _ in accounts]
            ba.append([CATEGORIES.get(cat, (cat,))[0]] + [float(v) if v else None for v in vals] + [float(sum(vals))])
        ba.column_dimensions["A"].width = 44
        for i in range(2, len(accounts) + 3):
            ba.column_dimensions[ba.cell(row=1, column=i).column_letter].width = 20

    cr = card_reconciliation(rows)
    if cr["cards"] or cr["unmatched_bank_payments"]:
        cc = wb.create_sheet("Credit cards")
        cc.append(["Credit card reconciliation"])
        cc["A1"].font = Font(bold=True, size=13)
        cc.append(["Card purchases are already counted as expenses under their categories. Bill payments to the card only settle those purchases, "
                   "so they are shown here for reconciliation and never counted as expenses."])
        cc.append([])
        cc.append(["Card", "Purchases (₹)", "No. of purchases", "Refunds (₹)", "Bill payments received (₹)", "Matched to a bank debit (₹)",
                   "Matched (no.)", "Not matched (₹)", "Not matched (no.)", "Paid in cash (₹)"])
        for c in cc[cc.max_row]:
            c.font, c.fill = bold, head_fill
        for c in cr["cards"]:
            cc.append([c["account"], float(c["purchases"]), c["purchase_count"], float(c["refunds"]), float(c["payments"]), float(c["matched"]),
                       c["matched_count"], float(c["unmatched"]), c["unmatched_count"], float(c["cash_payments"])])
        cc.append([])
        cc.append(["Total card bill payments this year (₹)", float(cr["total_paid"])])
        cc[cc.max_row][0].font = bold
        cc.append(["Compare with AIS → SFT 'Payment of credit card bills'. Issuers report payments of ₹10 lakh or more a year (₹1 lakh or more in cash)."])
        if cr["missing_statements"]:
            cc.append([])
            cc.append(["Missing card statements: bill payments seen in the bank with no matching card statement"])
            cc[cc.max_row][0].font = bold
            cc.append(["Card", "Issue", "Months", "Payments", "Amount (₹)", "What to do"])
            for c in cc[cc.max_row]:
                c.font, c.fill = bold, head_fill
            for g in cr["missing_statements"]:
                issue = {"not_identified": "Card not identified", "not_uploaded": "Statement not uploaded", "period_missing": "Month(s) missing", "not_matched": "Payment not found on statement"}[g["status"]]
                cc.append([g["card"], issue, ", ".join(g["months"]), g["count"], float(g["amount"]), g["message"]])
            cc.append([])
            cc.append(["Date", "Bank account", "Narration", "Amount (₹)", "Card"])
            for c in cc[cc.max_row]:
                c.font, c.fill = bold, head_fill
            for g in cr["missing_statements"]:
                for u in g["payments"]:
                    cc.append([u["date"], u["bank"], u["narration"], float(u["amount"]), g["card"]])
        for col, width in zip("ABCDEFGHIJ", [38, 16, 12, 14, 18, 18, 12, 16, 12, 14]):
            cc.column_dimensions[col].width = width

    fl = wb.create_sheet("Flags")
    fl.append(["Level", "Flag", "Detail"])
    for c in fl[1]:
        c.font, c.fill = bold, head_fill
    for f in purpose_flags(all_rows, purpose) + flags(rows, fy):
        fl.append([f["level"], f["title"], f["detail"]])
    for col, width in zip("ABC", [10, 60, 120]):
        fl.column_dimensions[col].width = width

    tx = wb.create_sheet("Transactions")
    show_purpose = business_mode or bool(purpose)
    tx.append(["FY", "Date", "Account", "Type", "Narration", "Debit (₹)", "Credit (₹)", "Balance (₹)", "Category", "Group", "Status", "Note", "Source file", "Source row"]
              + (["Purpose"] if show_purpose else []))
    for c in tx[1]:
        c.font, c.fill = bold, head_fill
    for t, a in rows:
        label, group, _ = CATEGORIES.get(t.category, (t.category, "review", ""))
        tx.append([t.financial_year, t.txn_date, a.name, a.kind, t.narration, float(t.debit), float(t.credit),
                   float(t.balance) if t.balance is not None else None, label, group.replace("_", " "), t.status.replace("_", " "), t.note or "", docs.get(t.document_id, ""), t.source_row]
                  + ([t.purpose] if show_purpose else []))
    for col, width in zip("ABCDEFGHIJKLMN", [9, 12, 22, 7, 60, 13, 13, 13, 30, 14, 20, 30, 30, 9]):
        tx.column_dimensions[col].width = width
    tx.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- closing the books: readiness checklist and ITR-head summary ----------

HEADS = [
    ("income", "Income", [
        ("salary", "Salary (take-home credits)", "Use gross salary and TDS from Form 16"),
        ("interest_income", "Interest", "Other sources; savings interest may qualify for 80TTA/80TTB"),
        ("dividend", "Dividends", "Other sources; match with AIS"),
        ("rental_income", "Rent received", "House property"),
        ("business_receipt", "Business / professional receipts", "Business or profession"),
        ("investment_redemption", "Investment sales / redemptions", "Capital gains: use the broker/AMC statement"),
        ("tax_refund", "Income-tax refund", "Not income; interest on it is"),
    ]),
    ("tax", "Tax paid", [
        ("tax_payment", "Income tax / TDS paid", "Match with 26AS challans"),
        ("gst_payment", "GST paid", "Reconcile with GSTR-3B"),
    ]),
    ("deductions", "Possible deductions", [
        ("investment", "Investments (MF, PPF, NPS, FD)", "80C / 80CCD, if eligible"),
        ("insurance", "Insurance premiums", "Life 80C · health 80D"),
        ("loan_emi", "Loan EMIs", "Home loan: interest 24(b), principal 80C"),
        ("rent", "Rent paid", "HRA / 80GG"),
        ("education", "Tuition fees", "80C, for children"),
        ("donation", "Donations", "80G, with receipts"),
    ]),
]


def final_heads(rows) -> list[dict]:
    by_cat: dict[str, list] = defaultdict(lambda: [Decimal(), 0])
    for t, _ in rows:
        amt = t.credit if t.credit > 0 else t.debit
        by_cat[t.category][0] += amt
        by_cat[t.category][1] += 1
    out = []
    for key, title, items in HEADS:
        lines = [{"category": c, "label": label, "hint": hint, "amount": str(by_cat[c][0]), "count": by_cat[c][1]}
                 for c, label, hint in items if by_cat[c][1]]
        out.append({"key": key, "title": title, "lines": lines, "total": str(sum((Decimal(l["amount"]) for l in lines), Decimal()))})
    # Spending: everything that is an expense, largest categories first.
    spend = defaultdict(lambda: [Decimal(), 0])
    for t, _ in rows:
        group = CATEGORIES.get(t.category, ("", "review", ""))[1]
        if group in {"expense"} and t.debit > 0:
            spend[t.category][0] += t.debit
            spend[t.category][1] += 1
    lines = sorted(({"category": c, "label": CATEGORIES[c][0], "hint": "", "amount": str(v[0]), "count": v[1]} for c, v in spend.items()),
                   key=lambda l: -Decimal(l["amount"]))
    out.append({"key": "spending", "title": "Spending", "lines": lines, "total": str(sum((Decimal(l["amount"]) for l in lines), Decimal()))})
    return out


def closing_checklist(db, workspace_id: int, fy: str | None, rows, business_mode: bool = False) -> list[dict]:
    """What still stands between these statements and final books, each with where to fix it."""
    items = []
    cov = statement_coverage(db, workspace_id, fy)
    gaps = []
    for acc in cov["accounts"]:
        present = [m for m in cov["months"] if acc["months"][m]["status"] != "missing"]
        if present:
            inside = cov["months"][cov["months"].index(present[0]):cov["months"].index(present[-1]) + 1]
            gaps += [f"{acc['account']} {m}" for m in inside if acc["months"][m]["status"] == "missing"]
    missing_cards = missing_card_statements(rows)
    n = len(gaps) + len(missing_cards)
    items.append({"key": "statements", "title": "Every statement uploaded", "done": n == 0, "count": n,
                  "detail": ("No gaps between uploaded statements, and every card bill paid from the bank has its card statement."
                             if n == 0 else
                             f"{len(gaps)} month(s) missing between uploaded statements" + (f"; {len(missing_cards)} card(s) with bill payments but no statement" if missing_cards else "") + "."),
                  "action": {"view": "coverage", "filter": "", "label": "Open statement coverage"}})
    review = [t for t, _ in rows if t.status == "needs_review"]
    items.append({"key": "categories", "title": "Every transaction categorised", "done": not review, "count": len(review),
                  "amount": str(sum((t.debit + t.credit for t in review), Decimal())),
                  "detail": "All transactions have a category." if not review else f"{len(review)} UPI/NEFT or unrecognised transaction(s) need a category.",
                  "action": {"view": "transactions", "filter": "needs_review", "label": "Categorise"}})
    open_items = [t for t, _ in rows if t.status in {"unmatched", "ambiguous"}]
    items.append({"key": "matching", "title": "Card bills and self-transfers matched", "done": not open_items, "count": len(open_items),
                  "detail": "Every card bill payment and self-transfer has its other side." if not open_items
                  else f"{len(open_items)} payment(s) have no other side yet: usually a statement that isn't uploaded.",
                  "action": {"view": "reconciliation", "filter": "", "label": "Open reconciliation"}})
    big = [t for t, _ in rows if t.credit >= 2 * LAKH and CATEGORIES.get(t.category, ("", "review", ""))[1] == "review" and not t.note]
    items.append({"key": "large_credits", "title": "Large credits explained", "done": not big, "count": len(big),
                  "amount": str(sum((t.credit for t in big), Decimal())),
                  "detail": "Every credit of ₹2 lakh or more has a category or a note." if not big
                  else f"{len(big)} credit(s) of ₹2 lakh or more have no category or note. The tax department can ask for their source.",
                  "action": {"view": "transactions", "filter": "needs_review", "label": "Explain them"}})
    if business_mode:
        unknown = [t for t, _ in rows if t.purpose == "unknown"]
        items.append({"key": "purpose", "title": "Business or personal decided", "done": not unknown, "count": len(unknown),
                      "detail": "Every transaction is marked business or personal." if not unknown
                      else f"{len(unknown)} transaction(s) are neither business nor personal yet and are left out of the business report.",
                      "action": {"view": "transactions", "filter": "purpose:unknown", "label": "Decide"}})
    return items
