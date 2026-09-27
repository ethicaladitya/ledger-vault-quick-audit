import io, zipfile
from tests.conftest import signup
from tests.test_parsers import HDFC_CSV

BANK = "date,narration,debit,credit\n2025-04-10,CREDIT CARD PAYMENT,35000,\n2025-04-12,UPI-RAMESH,500,\n"
CARD = "date,narration,debit,credit\n2025-04-11,PAYMENT RECEIVED THANK YOU,,35000\n2025-04-03,SWIGGY,1200,\n"


def upload(client, headers, files, account, kind="bank"):
    r = client.post("/imports/upload", headers=headers, data={"account_name": account, "kind": kind},
                    files=[("files", (name, content.encode() if isinstance(content, str) else content, "text/csv")) for name, content in files])
    assert r.status_code == 200, r.text
    return r.json()


def test_upload_review_report_and_export(client):
    h = signup(client)
    res = upload(client, h, [("bank.csv", BANK)], "HDFC Savings")
    assert res["files"][0]["transactions"] == 2
    res = upload(client, h, [("card.csv", CARD)], "HDFC Regalia", "card")
    assert res["reconciliation"]["confirmed"] == 1

    assert client.get("/years", headers=h).json()[0]["financial_year"] == "2025-26"
    dash = client.get("/dashboard?fy=2025-26", headers=h).json()
    assert dash["expenses"] == "1700.00"  # 500 UPI + 1200 card purchase; the 35,000 bill payment is neutral
    assert dash["neutral"] == "35000.00"
    assert dash["exceptions"] == 1

    review = client.get("/transactions?status=exceptions", headers=h).json()
    assert review["total"] == 1
    tx = review["items"][0]
    r = client.patch(f"/transactions/{tx['id']}", headers=h, json={"category": "rent", "note": "Flat rent to landlord"})
    assert r.status_code == 200 and r.json()["status"] == "ok"
    assert client.get("/dashboard", headers=h).json()["exceptions"] == 0

    rec = client.get("/reconciliation", headers=h).json()
    assert len(rec["matched"]) == 1 and not rec["open"]
    report = client.get("/report?fy=2025-26", headers=h).json()
    assert any(c["category"] == "rent" for c in report["categories"])
    x = client.get("/export.xlsx?fy=2025-26", headers=h)
    assert x.status_code == 200 and x.content[:2] == b"PK"


def test_zip_duplicate_overlap_and_delete(client):
    h = signup(client)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("statements/hdfc.csv", HDFC_CSV)
        z.writestr("__MACOSX/._hdfc.csv", "junk")
    res = upload(client, h, [("fy.zip", buf.getvalue())], "HDFC Savings")
    assert res["files"][0]["transactions"] == 3
    assert upload(client, h, [("again.csv", HDFC_CSV)], "HDFC Savings")["files"][0]["duplicate"] is True
    # A different file covering an overlapping period only adds the new row.
    overlap = HDFC_CSV.replace("STATEMENT SUMMARY", "06/04/25,NEW ROW,1,06/04/25,99.00,,1.00\nSTATEMENT SUMMARY")
    assert upload(client, h, [("overlap.csv", overlap)], "HDFC Savings")["files"][0]["transactions"] == 1
    docs = client.get("/documents", headers=h).json()
    assert len(docs) == 2
    assert client.delete(f"/documents/{docs[0]['id']}", headers=h).json()["removed"] == 1
    assert client.get("/transactions", headers=h).json()["total"] == 3


def test_bad_files_report_errors(client):
    h = signup(client)
    res = upload(client, h, [("statement.pdf", b"%PDF-1.4"), ("junk.csv", "hello,world\n1,2\n")], "X")
    assert "PDF" in res["files"][0]["error"]
    assert "header" in res["files"][1]["error"]


def test_workspaces_are_isolated(client):
    a = signup(client, "a@example.com")
    b = signup(client, "b@example.com")
    upload(client, a, [("bank.csv", BANK)], "A bank")
    # Same file in another workspace is not treated as a duplicate and stays private.
    assert upload(client, b, [("bank.csv", BANK)], "B bank")["files"][0]["transactions"] == 2
    tx_a = client.get("/transactions", headers=a).json()["items"]
    assert all(t["account"] == "A bank" for t in tx_a)
    assert client.patch(f"/transactions/{tx_a[0]['id']}", headers=b, json={"note": "x"}).status_code == 404
    doc_a = client.get("/documents", headers=a).json()[0]["id"]
    assert client.delete(f"/documents/{doc_a}", headers=b).status_code == 404
    for path in ["/transactions", "/reconciliation", "/report", "/dashboard", "/documents", "/export.xlsx"]:
        assert client.get(path).status_code == 401


def test_registration_closes_after_first_user(client, monkeypatch):
    from app import main
    monkeypatch.setattr(main, "ALLOW_REGISTRATION", False)
    assert client.get("/auth/config").json()["registration_open"] is True
    signup(client)
    assert client.get("/auth/config").json()["registration_open"] is False
    r = client.post("/auth/register", json={"email": "intruder@example.com", "password": "correct-horse-battery", "full_name": "Nope"})
    assert r.status_code == 403


def test_login_rate_limit(client):
    signup(client, "rl@example.com")
    for _ in range(10):
        assert client.post("/auth/login", json={"email": "rl@example.com", "password": "wrong-password"}).status_code == 401
    assert client.post("/auth/login", json={"email": "rl@example.com", "password": "correct-horse-battery"}).status_code == 429


DIVIDENDS = """date,narration,debit,credit
2025-06-21,ACH C- LARSEN AND TOUBRO LI-25213838,,34
2025-07-28,ACH C- HCL 2ND INTDIV25 26-601092,,24
2025-08-01,NEFT CR-GOATLIFE FARMS,,500
2025-08-02,NEFT CR-GOATLIFE FARMS REF 2,,700
2025-09-01,POS 416021XXXXXX9685 GOATLIFE,300,
"""


def test_dividends_and_learning_from_corrections(client):
    h = signup(client)
    upload(client, h, [("idfc.csv", DIVIDENDS)], "IDFC Savings")
    items = client.get("/transactions", headers=h).json()["items"]
    by = {t["narration"]: t for t in items}
    assert by["ACH C- LARSEN AND TOUBRO LI-25213838"]["category"] == "dividend"
    assert by["ACH C- HCL 2ND INTDIV25 26-601092"]["category"] == "dividend"
    first = by["NEFT CR-GOATLIFE FARMS"]
    r = client.patch(f"/transactions/{first['id']}", headers=h, json={"category": "business_receipt"}).json()
    assert r["similar"] == {"key": "goatlife farms", "count": 1}  # the debit-side POS row is not "similar"
    r = client.patch(f"/transactions/{first['id']}", headers=h, json={"category": "business_receipt", "apply_similar": True}).json()
    assert r["applied"] == 1
    # Remembered for the next upload.
    upload(client, h, [("idfc2.csv", "date,narration,debit,credit\n2025-10-01,NEFT CR-GOATLIFE FARMS X,,900\n")], "IDFC Savings")
    newest = client.get("/transactions?limit=1", headers=h).json()["items"][0]
    assert newest["category"] == "business_receipt" and newest["status"] == "ok"


def test_rules_upgrade_keeps_user_choices(client, db):
    from app.models import AppMeta, Transaction
    from app import migrate
    h = signup(client)
    upload(client, h, [("idfc.csv", DIVIDENDS)], "IDFC Savings")
    rows = db.query(Transaction).order_by(Transaction.id).all()
    rows[0].category = "uncategorized"          # as an old rule set would have left it
    rows[1].category, rows[1].category_source = "rent", "user"  # the user's explicit choice
    db.get(AppMeta, "rules_version").value = "0"
    db.commit()
    migrate.reclassify_if_rules_changed()
    db.expire_all()
    assert db.get(Transaction, rows[0].id).category == "dividend"
    assert db.get(Transaction, rows[1].id).category == "rent"


def test_review_and_confirm_after_upload(client):
    h = signup(client)
    res = upload(client, h, [("idfc.csv", DIVIDENDS)], "IDFC Savings")
    doc = res["files"][0]["document_id"]
    review = client.get(f"/imports/review?docs={doc}", headers=h).json()
    assert review["statements"][0]["account"] == "IDFC Savings" and review["statements"][0]["transactions"] == 5
    groups = review["groups"]
    assert groups[0]["needs_review"] is True  # undecided groups come first
    goat_in = next(g for g in groups if g["key"] == "goatlife farms")
    assert goat_in["count"] == 2 and goat_in["direction"] == "in"
    payload = [{"tx_ids": g["tx_ids"], "category": "business_receipt" if g is goat_in else g["category"],
                "remember_key": g["key"] if g is goat_in else None} for g in groups]
    r = client.post("/imports/confirm", headers=h, json={"groups": payload}).json()
    assert r["updated"] == 5 and r["remembered"] == 1
    assert client.get("/dashboard", headers=h).json()["exceptions"] == 0
    # Rename and re-type the detected account.
    acc = review["statements"][0]["account_id"]
    assert client.patch(f"/accounts/{acc}", headers=h, json={"name": "IDFC FIRST Savings", "kind": "bank"}).json()["name"] == "IDFC FIRST Savings"
    other = signup(client, "other@example.com")
    assert client.get(f"/imports/review?docs={doc}", headers=other).json() == {"statements": [], "groups": []}
    assert client.patch(f"/accounts/{acc}", headers=other, json={"name": "x"}).status_code == 404


def test_report_per_account(client):
    from openpyxl import load_workbook
    h = signup(client)
    upload(client, h, [("bank.csv", BANK)], "HDFC Savings")
    upload(client, h, [("card.csv", CARD)], "HDFC Regalia", "card")
    full = client.get("/report", headers=h).json()
    accounts = {a["account"]: a for a in full["accounts"]}
    assert accounts["HDFC Savings"]["outflow"] == "500.00" and accounts["HDFC Savings"]["neutral"] == "35000.00"
    assert accounts["HDFC Regalia"]["outflow"] == "1200.00" and accounts["HDFC Regalia"]["kind"] == "card"
    assert full["totals"]["outflow"] == "1700.00"  # overall report unchanged
    card_id = accounts["HDFC Regalia"]["account_id"]
    one = client.get(f"/report?account_id={card_id}", headers=h).json()
    assert one["totals"]["outflow"] == "1200.00" and {c["category"] for c in one["categories"]} == {"dining", "card_settlement"}
    wb = load_workbook(io.BytesIO(client.get("/export.xlsx", headers=h).content))
    assert wb.sheetnames == ["Summary", "By account", "Flags", "Transactions"]
    x = client.get(f"/export.xlsx?account_id={card_id}", headers=h)
    assert "HDFC-Regalia" in x.headers["content-disposition"]
    assert load_workbook(io.BytesIO(x.content))["Transactions"].max_row == 3  # header + 2 card rows
    other = signup(client, "other2@example.com")
    assert client.get(f"/report?account_id={card_id}", headers=other).json()["totals"]["outflow"] == "0"
