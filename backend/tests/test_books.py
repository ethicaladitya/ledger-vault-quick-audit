"""Books: one ledger across bank + card with card bill payments (both legs) and self-transfers left out."""
from decimal import Decimal
from app.services.rules import classify, category_fits
from tests.conftest import signup
from tests.test_api import upload

BANK = ("date,narration,debit,credit\n"
        "2025-04-01,NEFT ACME PVT LTD SALARY APR,,100000\n"
        "2025-04-10,BBPS/BILLDESK/HDFC0012345/REF 88,1700,\n"   # the card bill, but the narration never says "card"
        "2025-04-12,UPI-RAMESH-KIRANA,500,\n")
CARD = ("date,narration,debit,credit\n"
        "2025-04-03,SWIGGY BANGALORE,1200,\n"
        "2025-04-05,AMAZON PAY INDIA,800,\n"
        "2025-04-06,AMAZON PAY INDIA,,300\n"                        # merchant credit without the word "refund"
        "2025-04-11,BPPY CC PAYMENT DP0160142 (Ref# ST2601600830),,1700\n")


def test_books_merge_bank_and_card_without_bill_payment_legs(client):
    h = signup(client)
    upload(client, h, [("bank.csv", BANK)], "HDFC Savings")
    res = upload(client, h, [("card.csv", CARD)], "HDFC Regalia", "card")
    assert res["reconciliation"]["confirmed"] == 1   # BBPS debit ↔ card payment, found by amount and date

    b = client.get("/books?fy=2025-26", headers=h).json()
    got = [(e["date"], e["account"], e["out"], e["in"], e["refund"]) for e in b["entries"]]
    assert got == [
        ("2025-04-01", "HDFC Savings", "0.00", "100000.00", False),
        ("2025-04-03", "HDFC Regalia", "1200.00", "0.00", False),
        ("2025-04-05", "HDFC Regalia", "800.00", "0.00", False),
        ("2025-04-06", "HDFC Regalia", "0.00", "300.00", True),
        ("2025-04-12", "HDFC Savings", "500.00", "0.00", False),
    ]
    assert b["entries"][-1]["running"] == "97800.00"
    assert b["totals"] == {"money_in": "100000.00", "money_out": "2500.00", "refunds": "300.00", "net_spend": "2200.00",
                           "net": "97800.00", "count": 5}
    left_out = {e["label"]: (e["count"], e["matched"], e["debit"], e["credit"]) for e in b["excluded"]}
    assert left_out == {"Credit card bills paid from the bank": (1, 1, "1700.00", "0.00"),
                        "Bill payments received on the cards": (1, 1, "0.00", "1700.00")}
    assert not b["gaps"]

    # The report agrees: the refund reduces spending and is not income; the bill payment is not spending.
    t = client.get("/report?fy=2025-26", headers=h).json()["totals"]
    assert (t["inflow"], t["outflow"], t["refunds"], t["card_payments"]) == ("100000.00", "2200.00", "300.00", "1700.00")
    x = client.get("/export.xlsx?fy=2025-26", headers=h)
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames[:2] == ["Summary", "Books"]
    narrations = [r[3] for r in wb["Books"].iter_rows(min_row=5, values_only=True) if r[0] and not isinstance(r[0], str)]
    assert narrations == ["NEFT ACME PVT LTD SALARY APR", "SWIGGY BANGALORE", "AMAZON PAY INDIA", "AMAZON PAY INDIA", "UPI-RAMESH-KIRANA"]


def test_relabel_from_a_link_is_undone_when_the_other_statement_goes(client):
    h = signup(client)
    upload(client, h, [("bank.csv", BANK)], "HDFC Savings")
    upload(client, h, [("card.csv", CARD)], "HDFC Regalia", "card")
    card_doc = next(d for d in client.get("/documents", headers=h).json() if d["account"] == "HDFC Regalia")
    client.delete(f"/documents/{card_doc['id']}", headers=h)
    bbps = next(t for t in client.get("/transactions", headers=h).json()["items"] if t["narration"].startswith("BBPS"))
    assert (bbps["category"], bbps["status"], bbps["category_source"]) == ("bill_payment", "needs_review", "rule")
    # Unexplained, it stays in the books (as "needs review") rather than silently disappearing.
    assert any(e["narration"].startswith("BBPS") for e in client.get("/books", headers=h).json()["entries"])


def test_card_credits_are_never_income_and_card_debits_never_bill_payments():
    assert classify("BPPY CC PAYMENT DP0160142 (Ref# ST2601600830)", False, "card") == "card_settlement"
    assert classify("AMAZON PAY INDIA", False, "card") == "refund_reversal"
    assert classify("RAZORPAY SETTLEMENT", False, "card") == "refund_reversal"       # not a business receipt
    assert classify("NEFT ACME SALARY", False, "card") == "card_settlement"          # a payment by NEFT, not salary
    assert classify("BPPY CC PAYMENT DP0160142", True, "card") != "card_settlement"  # a misread sign stays visible
    # A payee rule taught on purchases doesn't turn a refund from that payee into income, or a salary rule a debit.
    assert classify("SWIGGY BANGALORE", False, "bank", {"swiggy bangalore": "dining"}) != "dining"
    assert classify("ACME PVT", True, "bank", {"acme pvt": "salary"}) != "salary"
    assert category_fits("dining", True) and not category_fits("dining", False) and not category_fits("salary", True)


def test_bill_payment_worded_like_a_refund_is_excluded_and_cards_break_down(client):
    h = signup(client)
    upload(client, h, [("bank.csv", "date,narration,debit,credit\n2025-05-10,UPI/DR/1/CRED/cred.club@axisb/Payment,2000,\n")], "HDFC Savings")
    card = ("date,narration,debit,credit\n"
            "2025-05-01,SWIGGY BANGALORE,1500,\n"
            "2025-05-02,AMAZON PAY INDIA,900,\n"
            "2025-05-04,AMAZON PAY INDIA - 12,,400\n"          # a real refund (reward points reversed)
            "2025-05-06,LATE FEE,100,\n"
            "2025-05-10,DREAMPLUG TECHNOLOGIES,,2000\n")     # the CRED payment arriving, worded like a merchant
    res = upload(client, h, [("card.csv", card)], "HDFC Regalia", "card")
    assert res["reconciliation"]["confirmed"] == 1
    b = client.get("/books", headers=h).json()
    assert [e["narration"] for e in b["entries"]] == ["SWIGGY BANGALORE", "AMAZON PAY INDIA", "AMAZON PAY INDIA - 12", "LATE FEE"]
    assert b["totals"]["refunds"] == "400.00"
    c = b["cards"]["cards"][0]
    assert (c["paid"], c["payments"], c["purchases"], c["charges"], c["refunds"], c["net_spend"]) == \
        ("2000.00", 1, "2400.00", "100.00", "400.00", "2100.00")
    assert [(k["category"], k["amount"]) for k in c["categories"]] == [("dining", "1500.00"), ("shopping", "900.00")]
    assert b["cards"]["paid_total"] == "2000.00" and b["cards"]["unassigned"]["count"] == 0


def test_statement_coverage_grid(client):
    h = signup(client)
    upload(client, h, [("bank.csv", BANK)], "HDFC Savings")
    upload(client, h, [("card-apr.csv", CARD)], "HDFC Regalia", "card")
    upload(client, h, [("card-jun.csv", "date,narration,debit,credit\n2025-05-20,ZOMATO,300,\n2025-06-18,UBER,200,\n")], "HDFC Regalia", "card")
    cov = client.get("/coverage?fy=2025-26", headers=h).json()
    assert cov["months"][0] == "2025-04" and cov["months"][-1] == "2026-03" and len(cov["months"]) == 12
    card = next(a for a in cov["accounts"] if a["account"] == "HDFC Regalia")
    marks = {m: c["status"] for m, c in card["months"].items()}
    assert marks["2025-04"] == "ok" and marks["2025-06"] == "ok"      # a statement sits in the month its cycle ends
    assert marks["2025-05"] == "missing" and card["uploaded"] == 2
    assert cov["accounts"][0]["account"] == "HDFC Regalia"               # cards first
    bank = next(a for a in cov["accounts"] if a["kind"] == "bank")
    assert bank["months"]["2025-04"]["status"] == "ok" and bank["months"]["2025-05"]["status"] == "missing"
    assert sum(cov["unexplained"].values()) == 0                          # the one bill payment matched
